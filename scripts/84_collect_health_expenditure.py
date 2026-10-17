# -*- coding: utf-8 -*-
"""Health expenditure for every panel country, 2001-2022 (the years of the population and migrant
series): current health expenditure (CHE) as a percentage of GDP.

Sources
  39 countries  World Bank WDI SH.XPD.CHEX.GD.ZS, which republishes the WHO Global Health
                Expenditure Database (GHED). Checked cell by cell against WHO's own API
                (GHO indicator GHED_CHEGDP_SHA2011); a value is used only if the two agree.
  Taiwan        Not in WDI or GHED. Ministry of Health and Welfare (MOHW), National Health
                Expenditure 2024 (中華民國113年國民醫療保健支出, published February 2026), Table 1,
                column "CHE占GDP比重 (%)": the same concept (current expenditure, i.e. without
                capital formation), on the latest national-accounts revision. Read from MOHW's
                official workbook (113年NHE統計表.xlsx, sheet 表1) and checked against the same
                table in the published report (PDF page 18); each percentage is also recomputed
                from the CHE and GDP columns.

Usage
  python scripts/84_collect_health_expenditure.py           parse the archived files only
  python scripts/84_collect_health_expenditure.py --fetch   download them first, and render the
                                                            mirrors, extracts and publisher pages

Writes data/extension_staging/health_expenditure_cells.csv in the collectors' cell format. The
file name does not match the api_*/national_* pattern on purpose: 74_assemble_panel.py, which
guards the 2001 extension, does not read it; 85_add_health_expenditure.py puts the cells into the
panel.
"""
import json
import os
import re
import sys
import urllib.parse

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import ext_lib as L                                    # noqa: E402

SITE = os.path.dirname(HERE)
D = os.path.join(SITE, 'data')
API = os.path.join(SITE, 'evidence', 'api')
PUB = os.path.join(API, 'publisher_pages')
TWN_DIR = os.path.join(SITE, 'evidence', 'countries', 'TWN')
COLLECTED = '2026-10-17'
YEARS = list(range(2001, 2023))
VAR = 'health_expenditure_pct_gdp'

WB_FILE = 'wb_SH_XPD_CHEX_GD_ZS_2001-2022_retrieved_%s.json' % COLLECTED
WB_URL = ('https://api.worldbank.org/v2/country/all/indicator/SH.XPD.CHEX.GD.ZS?format=json'
          '&date=2001:2022&per_page=20000')
WB_META = 'wb_SH_XPD_CHEX_GD_ZS_indicator_metadata_retrieved_%s.json' % COLLECTED
WB_META_URL = 'https://api.worldbank.org/v2/indicator/SH.XPD.CHEX.GD.ZS?format=json'
WHO_FILE = 'who_GHED_CHEGDP_SHA2011_2001-2022_retrieved_%s.json' % COLLECTED
WHO_URL = ('https://ghoapi.azureedge.net/api/GHED_CHEGDP_SHA2011?$filter='
           + urllib.parse.quote('TimeDim ge 2001 and TimeDim le 2022'))

TW_PAGE_URL = 'https://dep.mohw.gov.tw/DOS/lp-5071-113.html'   # MOHW statistics page listing both files
TW_XLSX_URL = 'https://www.mohw.gov.tw/dl-99630-35f4025b-4521-4ff4-84e8-22979a53a7c3.html'
TW_PDF_URL = 'https://www.mohw.gov.tw/dl-99631-cfe880cf-a90b-4eba-b240-351bf8abfd75.html'
TW_XLSX = 'health_expenditure__MOHW_NHE2024_tables__retrieved_%s.xlsx' % COLLECTED
TW_PDF = 'health_expenditure__MOHW_NHE2024_report__retrieved_%s.pdf' % COLLECTED
TW_HTML = 'health_expenditure__MOHW_DOS_NHE_page__retrieved_%s.html' % COLLECTED
TW_PAGE = 18                     # PDF page of Table 1 (printed page 12)
TW_EXTRACT = 'SNAPSHOT__health_expenditure__MOHW_NHE2024_table1_2001-2022'
TW_PAGESHOT = 'SNAPSHOT__health_expenditure__MOHW_DOS_NHE_page'

WB_NAME = ('World Bank, World Development Indicators (SH.XPD.CHEX.GD.ZS, current health '
           'expenditure % of GDP), sourced from the WHO Global Health Expenditure Database')
TW_NAME = ('Taiwan Ministry of Health and Welfare, Department of Statistics, National Health '
           'Expenditure 2024 (中華民國113年國民醫療保健支出, February 2026), Table 1 '
           '國民醫療保健支出指標, column CHE占GDP比重')
PUBLISHER_PAGES = [
    dict(key='worldbank_SH_XPD_CHEX_GD_ZS', publisher='World Bank',
         dataset='World Development Indicators - Current health expenditure (% of GDP) (SH.XPD.CHEX.GD.ZS)',
         page_url='https://data.worldbank.org/indicator/SH.XPD.CHEX.GD.ZS', backs_payload=WB_FILE),
    dict(key='who_ghed_che_pct_gdp', publisher='WHO',
         dataset='Global Health Expenditure Database - Current health expenditure (CHE) as percentage '
                 'of gross domestic product (GDP) (%) (GHED_CHEGDP_SHA2011)',
         page_url='https://www.who.int/data/gho/data/indicators/indicator-details/GHO/'
                  'current-health-expenditure-(che)-as-percentage-of-gross-domestic-product-(gdp)-(-)',
         backs_payload=WHO_FILE),
]

P = pd.read_csv(os.path.join(D, 'panel_final.csv'))
ISO = sorted(P.iso3.unique())
FETCH = '--fetch' in sys.argv


def fetch(url, dest, accept):
    code, size = L.download(url, dest, accept=accept)
    print('%-72s http %s %10s bytes' % (os.path.basename(dest)[:72], code, format(size, ',')))
    if code != '200' or size < 1000:
        sys.exit('download failed: ' + url)
    return size


def fetch_small(url, dest, tries=4):
    """The indicator-metadata response is under 1 kB, so fetch() would reject it; the World Bank
    endpoint also fails intermittently, so it is retried."""
    import time
    for _ in range(tries):
        code, size = L.download(url, dest, accept='application/json')
        try:
            if code == '200' and json.load(open(dest, encoding='utf-8-sig'))[1][0]['sourceNote']:
                print('%-72s http %s %10s bytes' % (os.path.basename(dest)[:72], code, format(size, ',')))
                return size
        except Exception:
            pass
        time.sleep(5)
    sys.exit('download failed: ' + url)


def upsert(path, rows, key):
    df = pd.read_csv(path)
    new = pd.DataFrame(rows)
    keep = df[~df[key].isin(new[key])]
    pd.concat([keep, new], ignore_index=True)[list(df.columns)].to_csv(path, index=False,
                                                                        encoding='utf-8-sig')


if FETCH:
    new = []
    for f, url, desc, pub in (
            (WB_FILE, WB_URL, 'Current health expenditure % of GDP (SH.XPD.CHEX.GD.ZS, from WHO '
                              'GHED), all countries 2001-2022', 'World Bank WDI'),
            (WHO_FILE, WHO_URL, 'WHO Global Health Expenditure Database via the GHO OData API, '
                                'current health expenditure % of GDP (GHED_CHEGDP_SHA2011), all '
                                'countries 2001-2022; the independent check of the World Bank '
                                'values', 'WHO'),
            (WB_META, WB_META_URL, 'Indicator metadata of SH.XPD.CHEX.GD.ZS: name, definition '
                                   '(sourceNote) and source organisation', 'World Bank WDI')):
        size = fetch(url, os.path.join(API, f), 'application/json') if f != WB_META else             fetch_small(url, os.path.join(API, f))
        json.load(open(os.path.join(API, f), encoding='utf-8-sig'))       # must parse
        new.append(dict(file=f, publisher=pub, description=desc, query_url=url, bytes=size,
                        path='evidence/api/' + f))
    upsert(os.path.join(D, 'api_snapshots.csv'), new, 'file')
    fetch(TW_XLSX_URL, os.path.join(TWN_DIR, TW_XLSX), 'application/octet-stream')
    fetch(TW_PDF_URL, os.path.join(TWN_DIR, TW_PDF), 'application/pdf')
    fetch(TW_PAGE_URL, os.path.join(TWN_DIR, TW_HTML), 'text/html')


# ------------------------------------------------------------------ World Bank and WHO
def load_json(rel):
    return json.load(open(os.path.join(API, rel), encoding='utf-8-sig'))


meta = load_json(WB_META)[1][0]
assert meta['id'] == 'SH.XPD.CHEX.GD.ZS' and 'capital health expenditures' in meta['sourceNote'], meta
wbj = load_json(WB_FILE)
wb_updated = wbj[0].get('lastupdated', '')
wb = {(r['countryiso3code'], int(r['date'])): float(r['value'])
      for r in wbj[1] if r.get('value') is not None}
who, who_dates = {}, set()
for r in load_json(WHO_FILE)['value']:
    if r.get('SpatialDimType') == 'COUNTRY' and r.get('NumericValue') is not None:
        who[(r['SpatialDim'], int(r['TimeDim']))] = float(r['NumericValue'])
        who_dates.add(str(r.get('Date', ''))[:10])
who_date = max(who_dates) if who_dates else ''

rows, compared, identical, mismatch = [], 0, 0, []
for iso in ISO:
    if iso == 'TWN':
        continue
    for y in YEARS:
        v, w = wb.get((iso, y)), who.get((iso, y))
        if v is None:
            print('   no World Bank value for %s %d' % (iso, y))
            continue
        if w is not None:
            compared += 1
            if abs(v - w) < 1e-6:
                identical += 1
            else:
                mismatch.append((iso, y, v, w))
                continue                      # used only where the two agree
        rows.append(L.cell(
            iso, y, VAR, v, 'WB_SH_XPD_CHEX_GD_ZS', WB_NAME,
            'https://api.worldbank.org/v2/country/all/indicator/SH.XPD.CHEX.GD.ZS',
            'evidence/api/' + WB_FILE, 'Year (annual flow, as published by WHO GHED)', 'A',
            'annual', note='', flag='',
            verification=('Read directly from the archived API response (evidence/api/%s, World '
                          'Bank data last updated %s); identical to full precision to the WHO GHO '
                          'API value (evidence/api/%s, WHO data dated %s). Retrieved %s.'
                          % (WB_FILE, wb_updated, WHO_FILE, who_date, COLLECTED)),
            collected_on=COLLECTED))
print('World Bank vs WHO GHO: %d cells compared, %d identical, %d different'
      % (compared, identical, len(mismatch)))
assert not mismatch, mismatch[:5]
assert compared == (len(ISO) - 1) * len(YEARS)

# ------------------------------------------------------------------ Taiwan, MOHW Table 1
import openpyxl                                          # noqa: E402
import fitz                                              # noqa: E402

ws = openpyxl.load_workbook(os.path.join(TWN_DIR, TW_XLSX), data_only=True)['表1']
assert str(ws['A1'].value).startswith('表1'), ws['A1'].value
hdr = ' '.join(str(c.value or '') for c in ws[5])
assert 'CHE占GDP比重' in hdr and 'NHE占GDP比重' in hdr, hdr
notes_x = ' '.join(str(r[0].value or '') for r in ws.iter_rows(min_row=42, max_row=50))
assert '民國96年(含)起含長期照顧服務支出' in notes_x
xl = {}
for r in ws.iter_rows(min_row=8, max_row=60, values_only=True):
    m = re.match(r'^(\d{2,3})年$', str(r[0] or ''))
    if m:
        xl[int(m.group(1)) + 1911] = dict(pop=r[1], gdp=r[2], nhe=r[4], nhe_pct=r[6], che=r[7],
                                          che_pct=r[9])
assert all(y in xl for y in YEARS), sorted(set(YEARS) - set(xl))

# the same table in the published report, page 18 (printed page 12)
doc = fitz.open(os.path.join(TWN_DIR, TW_PDF))
text = doc[TW_PAGE - 1].get_text()
doc.close()
assert '表1' in text and 'CHE占GDP' in text, 'Table 1 not on page %d' % TW_PAGE
NUM = r'(-?[\d,]+(?:\.\d+)?|\.\.\.)'
pdf = {}
for m in re.finditer(r'(\d{2,3})年\s+' + r'\s+'.join([NUM] * 11), text):
    f = [m.group(i) for i in range(2, 13)]
    pdf[int(m.group(1)) + 1911] = dict(gdp=f[1], che=f[6], che_pct=f[8])


def num(s):
    return float(str(s).replace(',', ''))


ext_rows = []
for y in YEARS:
    r = xl[y]
    pub = float(r['che_pct'])
    assert abs(num(pdf[y]['che_pct']) - pub) < 1e-9, (y, pdf[y], pub)          # workbook == report
    assert num(pdf[y]['che']) == r['che'] and num(pdf[y]['gdp']) == r['gdp'], y
    re_pct = r['che'] / r['gdp'] * 100
    assert abs(round(re_pct, 2) - pub) < 0.0051, (y, pub, re_pct)
    note = ('MOHW current health expenditure (經常性醫療保健支出, CHE) as %% of GDP: CHE %s million '
            'NT$ / nominal GDP %s million NT$ = %.4f%%, published to two decimals as %.2f%%. MOHW '
            'compiles CHE within the OECD System of Health Accounts framework; Table 1 states that '
            'it is based on the latest national-accounts revision, so it can differ from earlier '
            'editions. Table note 4: long-term care services are included from 民國96年 (2007); the '
            'English version of that note in the PDF says "Since 1997", which does not match the '
            'Chinese, and the Chinese is followed here.'
            % (format(r['che'], ','), format(r['gdp'], ','), re_pct, pub))
    flag = ''
    if y < 2007:
        flag = 'narrower scope: excludes long-term care services (included from 2007)'
    elif y == 2007:
        flag = 'break in time series: long-term care services included from 2007'
    rows.append(L.cell(
        'TWN', y, VAR, pub, 'TWN_MOHW_NHE2024', TW_NAME, TW_PAGE_URL, TW_XLSX,
        'Calendar year (annual flow)', 'A', 'annual', note=note, flag=flag,
        verification=('Read directly from the official MOHW workbook (%s, sheet 表1, downloaded '
                      'from %s); identical to Table 1 of the published report (%s, PDF page %d, '
                      'printed page 12, downloaded from %s); the percentage recomputed from the '
                      'CHE and GDP columns rounds to the published value. Retrieved %s.'
                      % (TW_XLSX, TW_XLSX_URL, TW_PDF, TW_PAGE, TW_PDF_URL, COLLECTED)),
        collected_on=COLLECTED))
    ext_rows.append(([u'%d (民國%d年)' % (y, y - 1911), format(r['gdp'], ','), format(r['nhe'], ','),
                      '%.2f' % r['nhe_pct'], format(r['che'], ','), '%.2f' % pub], True))
print('Taiwan: %d values read from the MOHW workbook, identical to the report'
      % sum(1 for r in rows if r['iso3'] == 'TWN'))

out = L.write_stage('health_expenditure_cells.csv', rows)
print('cells written: %d (%d countries x %d years)' % (len(out), len(ISO), len(YEARS)))
assert len(out) == len(ISO) * len(YEARS)

# the comparison of the two international sources, cell by cell, for the record
pd.DataFrame([dict(iso3=iso, year=y, world_bank=wb.get((iso, y)), who_gho=who.get((iso, y)),
                   identical=('yes' if abs(wb[(iso, y)] - who[(iso, y)]) < 1e-6 else 'no'))
              for iso in ISO if iso != 'TWN' for y in YEARS]).to_csv(
    os.path.join(D, 'health_expenditure_wb_vs_who.csv'), index=False, encoding='utf-8-sig')

# ------------------------------------------------------------------ mirrors, extracts, publisher pages
if FETCH:
    grid = {}
    for r in rows:
        if r['iso3'] != 'TWN':
            grid.setdefault(r['iso3'], {})[r['year']] = r['value']
    mirrors = []
    for rel, title, query, src in (
            (WB_FILE, 'World Bank WDI SH.XPD.CHEX.GD.ZS - current health expenditure (% of GDP)',
             'Query: api.worldbank.org country/all/indicator/SH.XPD.CHEX.GD.ZS, 2001-2022', wb),
            (WHO_FILE, 'WHO GHO GHED_CHEGDP_SHA2011 - current health expenditure (% of GDP)',
             'Query: ghoapi.azureedge.net/api/GHED_CHEGDP_SHA2011, TimeDim 2001-2022', who)):
        g = {iso: {y: src[(iso, y)] for y in YEARS if (iso, y) in src} for iso in grid}
        base = os.path.join(API, 'MIRROR__' + os.path.splitext(rel)[0])
        made = L.grid_mirror(title, '%s. Retrieved %s. %d panel countries; values shown to 2 '
                             'decimals, the payload holds full precision.' % (query, COLLECTED, len(g)),
                             g, YEARS, base, fmt=lambda v: '%.2f' % v)
        mirrors.append(dict(scope='API', folder='evidence/api', source_file=rel,
                            mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    mb = os.path.join(API, 'MIRROR__' + os.path.splitext(WB_META)[0])
    L.table_extract('World Bank indicator metadata - SH.XPD.CHEX.GD.ZS', ['Field', 'Value'],
                    [(['id', meta['id']], False), (['name', meta['name']], True),
                     (['sourceNote', meta['sourceNote']], True),
                     (['sourceOrganization', meta['sourceOrganization']], True),
                     (['source', meta['source']['value']], False)], mb,
                    'Archived response %s from %s, retrieved %s.' % (WB_META, WB_META_URL, COLLECTED),
                    note='The fields of the archived JSON, unchanged. Rendered %s.' % COLLECTED)
    mirrors.append(dict(scope='API', folder='evidence/api', source_file=WB_META,
                        mirrors='%s;%s' % (os.path.basename(mb + '.pdf'), os.path.basename(mb + '.png')),
                        n_mirrors=2))
    # the report: the page that carries Table 1
    pdf_page, png_page = L.page_extract(os.path.join(TWN_DIR, TW_PDF), TW_PAGE,
                                        os.path.join(TWN_DIR, 'MIRROR__' + os.path.splitext(TW_PDF)[0]))
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_PDF,
                        mirrors='%s;%s' % (os.path.basename(pdf_page), os.path.basename(png_page)),
                        n_mirrors=2))
    # the workbook: sheet 表1 as it stands, every row
    allrows = []
    for r in ws.iter_rows(min_row=8, max_row=41, values_only=True):
        allrows.append(([r[0]] + ['' if v is None else (('%.2f' % v) if isinstance(v, float) else
                                                         format(v, ',') if isinstance(v, int) else v)
                                  for v in r[1:12]],
                        bool(re.match(r'^(\d{2,3})年$', str(r[0] or '')) and
                             2001 <= int(str(r[0])[:-1]) + 1911 <= 2022)))
    xb = os.path.join(TWN_DIR, 'MIRROR__' + os.path.splitext(TW_XLSX)[0])
    L.table_extract('MOHW 113年NHE統計表.xlsx - sheet 表1 國民醫療保健支出指標 (all rows)',
                    ['年別', '年中人口數', 'GDP (百萬元)', 'GDP 年增率', 'NHE (百萬元)', 'NHE 年增率',
                     'NHE占GDP比重', 'CHE (百萬元)', 'CHE 年增率', 'CHE占GDP比重', '每人 NHE (元)',
                     '每人 NHE (美元)'], allrows, xb,
                    'Archived workbook %s, downloaded from %s on %s.' % (TW_XLSX, TW_XLSX_URL, COLLECTED),
                    note='Every row of sheet 表1 as stored in the workbook; the highlighted rows '
                         '(2001-2022) are the ones the panel uses (column CHE占GDP比重). Table note 4: '
                         '民國96年(含)起含長期照顧服務支出. Rendered %s.' % COLLECTED)
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_XLSX,
                        mirrors='%s;%s' % (os.path.basename(xb + '.pdf'), os.path.basename(xb + '.png')),
                        n_mirrors=2))
    # the statistics page that links both files
    # rendered from a copy that names the page's own address as its base, so its relative links
    # resolve to the MOHW site and not to the folder the copy sits in
    hb = os.path.join(TWN_DIR, 'MIRROR__' + os.path.splitext(TW_HTML)[0])
    raw = open(os.path.join(TWN_DIR, TW_HTML), encoding='utf-8', errors='replace').read()
    tmp = hb + '.tmp.html'
    open(tmp, 'w', encoding='utf-8').write(re.sub(r'(<head[^>]*>)', r'\1<base href="%s">' % TW_PAGE_URL,
                                                  raw, count=1, flags=re.I))
    made = L.chrome_render('file:///' + tmp.replace(chr(92), '/'), pdf=hb + '.pdf', png=hb + '.png')
    os.remove(tmp)
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_HTML,
                        mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    sm = os.path.join(D, 'source_mirrors.csv')
    old = pd.read_csv(sm)
    old = old[~(old.source_file.isin([m['source_file'] for m in mirrors]))]
    pd.concat([old, pd.DataFrame(mirrors)], ignore_index=True).to_csv(sm, index=False, encoding='utf-8-sig')

    # snapshots attached to the statistics-page URL, which is the panel's source URL for Taiwan
    epdf, epng, made = L.table_extract(
        'Taiwan - current health expenditure as % of GDP, 2001-2022 (MOHW Table 1)',
        ['Year', 'Nominal GDP (million NT$)', 'NHE (million NT$)', 'NHE % GDP',
         'CHE (million NT$)', 'CHE % GDP (used)'], ext_rows,
        os.path.join(TWN_DIR, TW_EXTRACT),
        'Source: %s. Workbook %s (sheet 表1) and report %s (PDF page %d), both linked from %s.'
        % (TW_NAME, TW_XLSX, TW_PDF, TW_PAGE, TW_PAGE_URL),
        note='Rows read from the archived workbook; the report prints the same figures. The last '
             'column is the value used. Long-term care services are included from 2007 (table note '
             '4). Extract prepared %s.' % COLLECTED)
    L.add_snapshot_row('TWN', VAR, TW_PAGE_URL, epdf, epng, TW_XLSX,
                       'Rows 2001-2022 of MOHW Table 1 (sheet 表1 of the workbook; the full page of '
                       'the report is MIRROR__%s.pdf/.png).' % os.path.splitext(TW_PDF)[0],
                       captured=COLLECTED)
    sp_pdf, sp_png = os.path.join(TWN_DIR, TW_PAGESHOT + '.pdf'), os.path.join(TWN_DIR, TW_PAGESHOT + '.png')
    L.chrome_render(TW_PAGE_URL, pdf=sp_pdf, png=sp_png)
    L.add_snapshot_row('TWN', VAR, TW_PAGE_URL, sp_pdf, sp_png, TW_HTML,
                       'The MOHW Department of Statistics page 國民醫療保健支出(NHE), which links the '
                       'workbook (113年NHE統計表.xlsx) and the report (113年國民醫療保健支出電子書).',
                       status='captured', captured=COLLECTED)

    # the publishers' own pages for the two international series
    pubrows = []
    for p in PUBLISHER_PAGES:
        b = os.path.join(PUB, p['key'])
        made = L.chrome_render(p['page_url'], pdf=b + '.pdf', png=b + '.png', wait_ms=20000)
        pubrows.append(dict(p, status='captured' if len(made) == 2 else 'partial',
                            pdf='evidence/api/publisher_pages/%s.pdf' % p['key'] if b + '.pdf' in made else '',
                            png='evidence/api/publisher_pages/%s.png' % p['key'] if b + '.png' in made else '',
                            pdf_bytes=os.path.getsize(b + '.pdf') if b + '.pdf' in made else 0,
                            png_bytes=os.path.getsize(b + '.png') if b + '.png' in made else 0,
                            captured=COLLECTED, note=''))
        print('publisher page %-30s %s' % (p['key'], [os.path.basename(m) for m in made]))
    upsert(os.path.join(D, 'api_publisher_snapshots.csv'), pubrows, 'key')
