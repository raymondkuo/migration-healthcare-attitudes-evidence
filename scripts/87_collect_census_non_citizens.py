# -*- coding: utf-8 -*-
"""Census counts of non-citizens for countries with no annual foreign-national series (issue #9, decided by the
authors on 2026-10-18).

The authors decided, from the ISSP 2021 item they analyse (people who "do not hold citizenship of [country]"), to
carry ONE census value in a separate column, non_citizens_census, kept out of foreign_nationals so that column stays
a register/annual series:

  Australia 2021   ABS, Census of Population and Housing: Cultural diversity data summary, 2021, Table 2
                   (country of birth of person by Australian citizenship), row "Total", column "Not an Australian
                   citizen". Census night 10 August 2021, inside Australia's ISSP 2021 fieldwork.

Not added (issue #10): South Africa's 2011 census count (1,692,242 "No", Census 2011 in Brief, Table 2.17). South
Africa's ISSP 2021 fieldwork was September-November 2021, ten years later; the count stays documented in the known
issues.

Usage
  python scripts/87_collect_census_non_citizens.py           parse the archived files only
  python scripts/87_collect_census_non_citizens.py --fetch   download them first, and render the mirrors

Writes data/extension_staging/census_non_citizens_cells.csv (the collectors' cell format); the name does not match
the api_*/national_* pattern, so 74_assemble_panel.py does not read it. 88_add_census_non_citizens.py puts the cell
into the panel.
"""
import os
import re
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import ext_lib as L                                    # noqa: E402

SITE = os.path.dirname(HERE)
D = os.path.join(SITE, 'data')
AUS = os.path.join(SITE, 'evidence', 'countries', 'AUS')
COLLECTED = '2026-10-18'
VAR = 'non_citizens_census'

PAGE_URL = 'https://www.abs.gov.au/statistics/people/people-and-communities/cultural-diversity-census/2021'
XLSX_URL = PAGE_URL + '/Cultural%20diversity%20data%20summary.xlsx'
DATE_URL = 'https://www.abs.gov.au/census/about-census'
XLSX = 'non_citizens_census__ABS_Census2021_cultural_diversity_summary__retrieved_%s.xlsx' % COLLECTED
PAGE = 'non_citizens_census__ABS_Census2021_cultural_diversity_page__retrieved_%s.html' % COLLECTED
DATE = 'non_citizens_census__ABS_about_census__retrieved_%s.html' % COLLECTED
EXTRACT = 'SNAPSHOT__non_citizens_census__ABS_Census2021_Table2_total'
NAME = ('Australian Bureau of Statistics, Census of Population and Housing: Cultural diversity data summary, 2021, '
        'Table 2 (country of birth of person by Australian citizenship), row Total, column Not an Australian citizen')
FETCH = '--fetch' in sys.argv

if FETCH:
    for url, f, accept in ((XLSX_URL, XLSX, 'application/octet-stream'), (PAGE_URL, PAGE, 'text/html'),
                           (DATE_URL, DATE, 'text/html')):
        code, size = L.download(url, os.path.join(AUS, f), accept=accept)
        print('%-80s http %s %9s bytes' % (f[:80], code, format(size, ',')))
        if code != '200' or size < 1000:
            sys.exit('download failed: ' + url)

# ------------------------------------------------------------------ the count, from the official workbook
import openpyxl                                          # noqa: E402

wb = openpyxl.load_workbook(os.path.join(AUS, XLSX), data_only=True)
ws = wb['Table 2']
title = ' '.join(str(c.value or '') for row in ws.iter_rows(min_row=1, max_row=6) for c in row)
assert 'Cultural diversity data summary, 2021' in title and 'COUNTRY OF BIRTH OF PERSON BY AUSTRALIAN CITIZENSHIP' in title
head = [str(c.value or '').strip() for c in ws[9]][:5]
assert head == ['', 'Australian citizen', 'Australian citizen', 'Not stated', 'Total'], head
sub = str(ws.cell(8, 3).value or '').strip()
assert sub == 'Not an', sub                          # column C reads "Not an / Australian citizen"
foot = ' '.join(str(r[0].value or '') for r in ws.iter_rows(min_row=330, max_row=ws.max_row))
assert 'place of usual residence' in foot and 'Excludes overseas visitors' in foot
rows = [r for r in ws.iter_rows(min_row=10, max_row=ws.max_row, values_only=True) if r[0] == 'Total']
total = rows[-1]                                     # the last "Total" row is the national total (the first is Africa's)
citizen, non_citizen, not_stated, persons = (int(x) for x in total[1:5])
print('Australia 2021 census: citizens %s, not citizens %s, not stated %s, total %s'
      % tuple(format(x, ',') for x in (citizen, non_citizen, not_stated, persons)))
assert non_citizen == 2808214 and persons == 25422788
date_html = open(os.path.join(AUS, DATE), encoding='utf-8', errors='replace').read()
assert 'The 2021 Census was held on Tuesday 10 August 2021' in date_html

note = ('Census count, not an annual register: persons counted in the 2021 Census of Population and Housing who are not '
        'Australian citizens, by place of usual residence, excluding overseas visitors. Of the %s persons in the table, '
        '%s (%.1f%%) did not state their citizenship and are not in this count; ABS adds small random adjustments to every '
        'cell, so the components need not sum to the total exactly. Census night was 10 August 2021. Kept apart from '
        'foreign_nationals (an annual series elsewhere); the share in non_citizens_census_pct_pop uses the World Bank '
        'population of the panel, like every other share (on the census persons it would be %.1f%%).'
        % (format(persons, ','), format(not_stated, ','), not_stated / persons * 100, non_citizen / persons * 100))
cells = [L.cell('AUS', 2021, VAR, non_citizen, 'ABS_CENSUS2021_CD_T2', NAME, PAGE_URL, XLSX,
                'Census night, 10 August 2021', 'A', 'census', note=note,
                flag='census count (single year); not stated excluded',
                verification=('Read directly from the official ABS workbook (%s, sheet Table 2, national Total row; '
                              'downloaded from %s); census date from the archived ABS page %s. Retrieved %s.'
                              % (XLSX, XLSX_URL, DATE, COLLECTED)),
                collected_on=COLLECTED)]
L.write_stage('census_non_citizens_cells.csv', cells)
print('cells written: %d' % len(cells))

# ------------------------------------------------------------------ mirrors and the extract (only when fetching)
if FETCH:
    summary = [r for r in ws.iter_rows(min_row=320, max_row=ws.max_row, values_only=True)
               if r[0] in ('Total', 'Not stated', 'Inadequately described', 'At sea')]
    allrows = [([r[0]] + [format(int(v), ',') for v in r[1:5]], i == len(summary) - 1)
               for i, r in enumerate(summary)]
    base = os.path.join(AUS, 'MIRROR__' + os.path.splitext(XLSX)[0])
    L.table_extract('ABS Cultural diversity data summary, 2021 - Table 2, summary rows',
                    ['Country of birth', 'Australian citizen', 'Not an Australian citizen', 'Not stated', 'Total'],
                    allrows, base,
                    'Archived workbook %s, downloaded from %s on %s. Footnotes: based on place of usual residence; '
                    'excludes overseas visitors; small random adjustments are made to all cell values.'
                    % (XLSX, XLSX_URL, COLLECTED),
                    note='The rows below the country list of Table 2 as stored in the workbook; the highlighted national '
                         'total is the row read. Rendered %s.' % COLLECTED)
    mirrors = [dict(scope='AUS', folder='evidence/countries/AUS', source_file=XLSX,
                    mirrors='%s.pdf;%s.png' % (os.path.basename(base), os.path.basename(base)), n_mirrors=2)]
    for f, url in ((PAGE, PAGE_URL), (DATE, DATE_URL)):
        hb = os.path.join(AUS, 'MIRROR__' + os.path.splitext(f)[0])
        raw = open(os.path.join(AUS, f), encoding='utf-8', errors='replace').read()
        tmp = hb + '.tmp.html'
        open(tmp, 'w', encoding='utf-8').write(re.sub(r'(<head[^>]*>)', r'\1<base href="%s">' % url, raw, count=1,
                                                      flags=re.I))
        made = L.chrome_render('file:///' + tmp.replace(chr(92), '/'), pdf=hb + '.pdf', png=hb + '.png')
        os.remove(tmp)
        mirrors.append(dict(scope='AUS', folder='evidence/countries/AUS', source_file=f,
                            mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    sm = os.path.join(D, 'source_mirrors.csv')
    old = pd.read_csv(sm)
    old = old[~old.source_file.isin([m['source_file'] for m in mirrors])]
    pd.concat([old, pd.DataFrame(mirrors)], ignore_index=True).to_csv(sm, index=False, encoding='utf-8-sig')
    epdf, epng, made = L.table_extract(
        'Australia - persons who are not Australian citizens, 2021 Census',
        ['Item', 'Value'],
        [(['Not an Australian citizen (used)', format(non_citizen, ',')], True),
         (['Australian citizen', format(citizen, ',')], False),
         (['Not stated', format(not_stated, ',')], False),
         (['Total persons (place of usual residence, excl. overseas visitors)', format(persons, ',')], False),
         (['Census night', '10 August 2021'], False)],
        os.path.join(AUS, EXTRACT),
        'Source: %s. Workbook %s, linked from %s.' % (NAME, XLSX, PAGE_URL),
        note='Read from the national Total row of Table 2 in the archived workbook; census date from %s. Extract '
             'prepared %s.' % (DATE_URL, COLLECTED))
    L.add_snapshot_row('AUS', VAR, PAGE_URL, epdf, epng, XLSX,
                       'National total row of Table 2 (the full summary rows are in MIRROR__%s.pdf/.png).'
                       % os.path.splitext(XLSX)[0], captured=COLLECTED)
