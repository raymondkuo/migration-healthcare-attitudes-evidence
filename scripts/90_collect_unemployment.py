# -*- coding: utf-8 -*-
"""Unemployment rate for every panel country, 2001-2022 (the years of the population and migrant
series): unemployed persons as a percentage of the labour force, ages 15 and over.

Sources
  39 countries  ILO modelled estimates (ILOEST), November 2025 edition, indicator UNE_2EAP_SEX_AGE_RT
                (total, age 15+), read from the ILO's own SDMX API. Each value carries the ILO's
                observation status: "R" (Real value) where the estimate rests on national data;
                values without it are ILO imputations, which the ILO describes as not based on
                national data and not meant for country comparisons. Those are kept and flagged.
                Checked cell by cell against the World Bank's copy (WDI SL.UEM.TOTL.ZS, "modeled
                ILO estimate"), which carries more digits: a value is used only if the World Bank
                value rounds to it at the ILO's three decimals.
  Taiwan        Directorate-General of Budget, Accounting and Statistics (DGBAS), Yearbook of
                Manpower Survey Statistics 2025 (114年人力資源調查統計年報, March 2026), Table 2
                歷年人力資源調查重要指標, column 失業率 總計: the annual average of the monthly
                Manpower Survey, civilian population aged 15 and over. Read from DGBAS's official
                workbook (table2.xlsx) and checked against the same table in the published yearbook
                (PDF pages 20-21). The ILO also holds Taiwan, but imputes 2021 and 2022 although
                DGBAS published them, so the national series is used for every year and the ILO's
                values are kept beside it for comparison.

Usage
  python scripts/90_collect_unemployment.py           parse the archived files only
  python scripts/90_collect_unemployment.py --fetch   download them first, and render the
                                                      mirrors, extracts and publisher page

Writes data/extension_staging/unemployment_cells.csv in the collectors' cell format (not read by
74_assemble_panel.py; 91_add_unemployment.py puts the cells into the panel) and
data/unemployment_source_comparison.csv (ILO, World Bank and, for Taiwan, DGBAS side by side).

The ILOSTAT web pages (ilostat.ilo.org) answer automated requests with a bot check, so no page of
theirs is captured; the ILO's own description of the dataset, its edition date and the meaning of
its observation-status codes are archived from the SDMX API instead.
"""
import json
import os
import re
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import ext_lib as L                                    # noqa: E402

SITE = os.path.dirname(HERE)
D = os.path.join(SITE, 'data')
API = os.path.join(SITE, 'evidence', 'api')
PUB = os.path.join(API, 'publisher_pages')
TWN_DIR = os.path.join(SITE, 'evidence', 'countries', 'TWN')
COLLECTED = '2026-10-21'
YEARS = list(range(2001, 2023))
VAR = 'unemployment_rate'

ILO_BASE = 'https://sdmx.ilo.org/rest/data/ILO,DF_UNE_2EAP_SEX_AGE_RT/.A..SEX_T.AGE_YTHADULT_YGE15'
ILO_URL = ILO_BASE + '?startPeriod=2001&endPeriod=2022'
ILO_ACCEPT = 'application/vnd.sdmx.data+csv;version=1.0.0'
ILO_FILE = 'ilo_UNE_2EAP_SEX_AGE_RT_2001-2022_retrieved_%s.csv' % COLLECTED
ILO_DF_URL = 'https://sdmx.ilo.org/rest/dataflow/ILO/DF_UNE_2EAP_SEX_AGE_RT/1.0?detail=full'
ILO_DF = 'ilo_DF_UNE_2EAP_SEX_AGE_RT_dataflow_retrieved_%s.xml' % COLLECTED
ILO_CL_URL = 'https://sdmx.ilo.org/rest/codelist/ILO/CL_OBS_STATUS'
ILO_CL = 'ilo_CL_OBS_STATUS_codelist_retrieved_%s.xml' % COLLECTED
WB_FILE = 'wb_SL_UEM_TOTL_ZS_2001-2022_retrieved_%s.json' % COLLECTED
WB_URL = ('https://api.worldbank.org/v2/country/all/indicator/SL.UEM.TOTL.ZS?format=json'
          '&date=2001:2022&per_page=20000')
WB_META = 'wb_SL_UEM_TOTL_ZS_indicator_metadata_retrieved_%s.json' % COLLECTED
WB_META_URL = 'https://api.worldbank.org/v2/indicator/SL.UEM.TOTL.ZS?format=json'

TW_PAGE_URL = 'https://www.stat.gov.tw/News_Content.aspx?n=4001&s=236078'   # DGBAS page listing the yearbook files
TW_REL = 'https://ws.dgbas.gov.tw/001/Upload/463/relfile/11516/236078/'
TW_XLSX_URL = TW_REL + 'table2.xlsx'
TW_PDF_URL = TW_REL + 'f194afc4-ad75-46bf-aa65-02239bdfdf8d.pdf'
TW_XLSX = 'unemployment_rate__DGBAS_Manpower_Yearbook2025_table2__retrieved_%s.xlsx' % COLLECTED
TW_PDF = 'unemployment_rate__DGBAS_Manpower_Yearbook2025__retrieved_%s.pdf' % COLLECTED
TW_HTML = 'unemployment_rate__DGBAS_Manpower_Yearbook2025_page__retrieved_%s.html' % COLLECTED
TW_PAGES = (20, 21)              # PDF pages of Table 2: years and first columns, then the rates (printed 4-5)
TW_DEF_PAGE = 322                # 名詞與定義: labour force, unemployed, unemployment rate (printed page 292)
TW_EXTRACT = 'SNAPSHOT__unemployment_rate__DGBAS_Yearbook2025_table2_2001-2022'
TW_PAGESHOT = 'SNAPSHOT__unemployment_rate__DGBAS_Yearbook2025_page'   # not made: see the end of the script

ILO_NAME = ('International Labour Organization, ILO modelled estimates (ILOEST), November 2025: unemployment '
            'rate, total, age 15+ (UNE_2EAP_SEX_AGE_RT), ILO SDMX API')
TW_NAME = ('Taiwan Directorate-General of Budget, Accounting and Statistics, Yearbook of Manpower Survey '
           'Statistics 2025 (114年人力資源調查統計年報, March 2026), Table 2 歷年人力資源調查重要指標, '
           'column 失業率 總計')
PUBLISHER_PAGES = [
    dict(key='worldbank_SL_UEM_TOTL_ZS', publisher='World Bank',
         dataset='World Development Indicators - Unemployment, total (% of total labor force) (modeled ILO '
                 'estimate) (SL.UEM.TOTL.ZS)',
         page_url='https://data.worldbank.org/indicator/SL.UEM.TOTL.ZS', backs_payload=WB_FILE),
]

P = pd.read_csv(os.path.join(D, 'panel_final.csv'), usecols=['iso3', 'year'])
ISO = sorted(P.iso3.unique())
FETCH = '--fetch' in sys.argv


def fetch(url, dest, accept, minimum=1000):
    code, size = L.download(url, dest, accept=accept)
    print('%-72s http %s %10s bytes' % (os.path.basename(dest)[:72], code, format(size, ',')))
    if code != '200' or size < minimum:
        sys.exit('download failed: ' + url)
    return size


def fetch_small(url, dest, tries=4):
    """The indicator-metadata response is under 1 kB and the World Bank endpoint fails intermittently."""
    import time
    for _ in range(tries):
        code, size = L.download(url, dest, accept='application/json')
        try:
            if code == '200' and json.load(open(dest, encoding='utf-8-sig'))[1][0]['sourceNote']:
                print('%-72s http %s %10s bytes' % (os.path.basename(dest)[:72], code, format(size, ',')))
                return size
        except Exception:
            pass
        time.sleep(5)
    sys.exit('download failed: ' + url)


def upsert(path, rows, key):
    df = pd.read_csv(path)
    new = pd.DataFrame(rows)
    keep = df[~df[key].isin(new[key])]
    pd.concat([keep, new], ignore_index=True)[list(df.columns)].to_csv(path, index=False,
                                                                        encoding='utf-8-sig')


if FETCH:
    new = []
    for f, url, desc, pub, acc in (
            (ILO_FILE, ILO_URL, 'ILO modelled estimates (Nov. 2025), unemployment rate, total, age 15+ '
                                '(UNE_2EAP_SEX_AGE_RT), all countries 2001-2022, SDMX-CSV with observation '
                                'status and bounds', 'ILO', ILO_ACCEPT),
            (ILO_DF, ILO_DF_URL, 'SDMX dataflow DF_UNE_2EAP_SEX_AGE_RT: the ILO\'s name, edition, last-update '
                                 'date and description of the dataset (incl. its note on imputed observations)',
             'ILO', 'application/xml'),
            (ILO_CL, ILO_CL_URL, 'SDMX code list CL_OBS_STATUS: the meaning of the ILO\'s observation-status '
                                 'codes (R = Real value)', 'ILO', 'application/xml'),
            (WB_FILE, WB_URL, 'Unemployment, total (% of total labor force) (modeled ILO estimate) '
                              '(SL.UEM.TOTL.ZS), all countries 2001-2022; the independent check of the ILO '
                              'values', 'World Bank WDI', 'application/json'),
            (WB_META, WB_META_URL, 'Indicator metadata of SL.UEM.TOTL.ZS: name, definition (sourceNote) and '
                                   'source organisation (ILOEST)', 'World Bank WDI', 'application/json')):
        dest = os.path.join(API, f)
        size = fetch_small(url, dest) if f == WB_META else fetch(url, dest, acc)
        if f.endswith('.json'):
            json.load(open(dest, encoding='utf-8-sig'))                 # must parse
        new.append(dict(file=f, publisher=pub, description=desc, query_url=url, bytes=size,
                        path='evidence/api/' + f))
    upsert(os.path.join(D, 'api_snapshots.csv'), new, 'file')
    fetch(TW_XLSX_URL, os.path.join(TWN_DIR, TW_XLSX), 'application/octet-stream')
    fetch(TW_PDF_URL, os.path.join(TWN_DIR, TW_PDF), 'application/pdf')
    fetch(TW_PAGE_URL, os.path.join(TWN_DIR, TW_HTML), 'text/html')


# ------------------------------------------------------------------ the ILO's own metadata
def xml_text(rel):
    return re.sub(r'\s+', ' ', open(os.path.join(API, rel), encoding='utf-8').read())


dft = xml_text(ILO_DF)
ilo_name = re.search(r'<common:Name xml:lang="en">([^<]+)</common:Name>', dft).group(1)
assert 'ILO modelled estimates, Nov. 2025' in ilo_name, ilo_name
m = re.search(r'<common:AnnotationTitle>(\d\d)/(\d\d)/(\d{4})[^<]*</common:AnnotationTitle>\s*'
              r'<common:AnnotationType>LAST_UPDATE', dft)
ilo_updated = '%s-%s-%s' % (m.group(3), m.group(2), m.group(1))          # day/month/year in the response
ilo_desc = re.sub(r'<[^>]+>', '', re.search(r'<common:Description xml:lang="en">(.*?)</common:Description>',
                                            dft).group(1).replace('&lt;', '<').replace('&gt;', '>')
                  .replace('&amp;', '&')).replace('&nbsp;', ' ').strip()
assert 'Imputed observations are not based on national data' in ilo_desc, ilo_desc[:200]
clt = xml_text(ILO_CL)
codes = {mm.group(1): re.findall(r'<common:Name xml:lang="en">([^<]+)', mm.group(0))[0]
         for mm in re.finditer(r'<structure:Code id="([^"]+)".*?</structure:Code>', clt)}
assert codes.get('R') == 'Real value', codes

# ------------------------------------------------------------------ ILO and World Bank
ilo_df = pd.read_csv(os.path.join(API, ILO_FILE), dtype=str, keep_default_na=False)
assert set(ilo_df.SEX) == {'SEX_T'} and set(ilo_df.AGE) == {'AGE_YTHADULT_YGE15'}, 'unexpected series'
assert set(ilo_df.SOURCE) == {'ILO - Modelled Estimates'}
ilo = {(r.REF_AREA, int(r.TIME_PERIOD)): r for r in ilo_df.itertuples()}
meta = json.load(open(os.path.join(API, WB_META), encoding='utf-8-sig'))[1][0]
assert meta['id'] == 'SL.UEM.TOTL.ZS' and 'ILOEST' in meta['sourceOrganization'], meta
wbj = json.load(open(os.path.join(API, WB_FILE), encoding='utf-8-sig'))
wb_updated = wbj[0].get('lastupdated', '')
wb = {(r['countryiso3code'], int(r['date'])): float(r['value'])
      for r in wbj[1] if r.get('value') is not None}


def imputed(r):
    return r.OBS_STATUS != 'R'


rows, compared, agree, mismatch = [], 0, 0, []
for iso in ISO:
    if iso == 'TWN':
        continue
    for y in YEARS:
        r = ilo.get((iso, y))
        if r is None:
            sys.exit('no ILO value for %s %d' % (iso, y))
        v = float(r.OBS_VALUE)
        assert r.OBS_STATUS in ('R', ''), (iso, y, r.OBS_STATUS)
        if r.OBS_STATUS == 'R':
            assert r.UPPER_BOUND == r.LOWER_BOUND == r.OBS_VALUE, (iso, y)   # a real value has no interval
        w = wb.get((iso, y))
        compared += 1
        if w is not None and round(w, 3) == v:
            agree += 1
        else:
            mismatch.append((iso, y, v, w))
            continue                          # used only where the two agree
        if imputed(r):
            # one wording for every imputed value; the bounds of each are in the raw response and in
            # data/unemployment_source_comparison.csv
            note = ('ILO imputed value: the ILO gives it no observation status (its real values are marked R, '
                    '"Real value") and publishes a lower and an upper bound with it (in the raw response and in '
                    'data/unemployment_source_comparison.csv). The ILO describes imputed observations as not based '
                    'on national data, subject to high uncertainty and not to be used for country comparisons or '
                    'rankings (dataflow description, evidence/api/%s).' % ILO_DF)
            flag, stype = 'ILO imputed: not based on national data', 'model_estimate'
        else:
            note, flag, stype = ('ILO observation status R (Real value): the estimate rests on national data.',
                                 '', 'annual')
        rows.append(L.cell(
            iso, y, VAR, v, 'ILO_UNE_2EAP_SEX_AGE_RT', ILO_NAME, ILO_BASE, 'evidence/api/' + ILO_FILE,
            'Year (annual average, as estimated by the ILO)', 'A', stype, note=note, flag=flag,
            verification=('Read directly from the archived ILO SDMX response (evidence/api/%s; %s, last updated '
                          '%s); the World Bank WDI value (evidence/api/%s, data last updated %s) rounds to it at the '
                          'ILO\'s three decimals. Retrieved %s.'
                          % (ILO_FILE, ilo_name, ilo_updated, WB_FILE, wb_updated, COLLECTED)),
            collected_on=COLLECTED))
print('ILO vs World Bank: %d cells compared, %d agree at three decimals, %d differ'
      % (compared, agree, len(mismatch)))
assert not mismatch, mismatch[:5]
assert compared == (len(ISO) - 1) * len(YEARS)
n_imp = sum(1 for r in rows if r['flag'])
print('ILO imputed values among them: %d' % n_imp)

# ------------------------------------------------------------------ Taiwan, DGBAS Table 2
import openpyxl                                          # noqa: E402
import fitz                                              # noqa: E402

ws = openpyxl.load_workbook(os.path.join(TWN_DIR, TW_XLSX), data_only=True).worksheets[0]
cells_ = [[c for c in r] for r in ws.iter_rows(values_only=True)]
assert any('歷年人力資源調查重要指標' in str(v) for v in cells_[1] if v), cells_[1]
hdr = ' '.join(str(v) for r in cells_[5:11] for v in r if v)
assert '失業率' in hdr and 'Unemployment rate' in hdr and '單位：％' in ' '.join(str(v) for v in cells_[4] if v)
YR = re.compile(r'^\s*([０-９\d]{2,3})年平均 Ave\., (\d{4})\s*$')
xl, allrows = {}, []
for r in cells_[11:]:
    vals = [v for v in r if v is not None]
    if not vals:
        continue
    m = YR.match(str(vals[0]))
    if m:
        nums = vals[1:12]
        assert len(nums) == 11, vals
        xl[int(m.group(2))] = nums
    allrows.append(vals)
assert all(y in xl for y in YEARS), sorted(set(YEARS) - set(xl))

# the same table in the published yearbook: page 20 holds the year labels and the first five columns, page 21
# the remaining six (employed/population, employed/15+, employed/labour force, unemployment total, male, female)
doc = fitz.open(os.path.join(TWN_DIR, TW_PDF))
p20, p21 = doc[TW_PAGES[0] - 1].get_text(), doc[TW_PAGES[1] - 1].get_text()
pdef = re.sub(r'\s+', '', doc[TW_DEF_PAGE - 1].get_text())
doc.close()
assert '(六)失業率：失業者占勞動力之比率' in pdef and '(三)失業者：指在資料標準週內年滿15歲' in pdef, 'definitions'
labels = re.findall(r'([０-９\d]{2,3})年平均 Ave\., (\d{4})', p20)
pdf_years = [int(b) for _, b in labels]
right = re.findall(r'-?\d+\.\d\d', p21)
pdf = {y: right[6 * i + 3] for i, y in enumerate(pdf_years)}     # 4th of the six columns = unemployment, total

ext_rows = []
for y in YEARS:
    n = xl[y]
    pub = float(n[8])
    assert pdf[y] == '%.2f' % pub, (y, pdf[y], pub)                # workbook == yearbook
    assert abs(100 - float(n[7]) - pub) < 0.0051, (y, n[7], pub)    # employed share of the labour force + rate = 100
    r = ilo[('TWN', y)]
    tw_real = [yy for yy in YEARS if ilo[('TWN', yy)].OBS_STATUS == 'R']
    tw_imp = [yy for yy in YEARS if ilo[('TWN', yy)].OBS_STATUS != 'R']
    note = ('DGBAS Manpower Survey (人力資源調查): unemployed persons as a share of the labour force, civilian '
            'population aged 15 and over (excludes the armed forces and the institutional population), annual '
            'average of the monthly survey; published to two decimals (definitions: yearbook PDF page %d). The ILO '
            'modelled estimates hold Taiwan as well (data/unemployment_source_comparison.csv): status R (Real value) '
            'for %d-%d, imputed for %s although DGBAS published those years; the national series is used for every '
            'year.' % (TW_DEF_PAGE, tw_real[0], tw_real[-1], ' and '.join(map(str, tw_imp))))
    rows.append(L.cell(
        'TWN', y, VAR, pub, 'TWN_DGBAS_MANPOWER_2025', TW_NAME, TW_PAGE_URL, TW_XLSX,
        'Calendar year (annual average of the monthly survey)', 'A', 'annual', note=note, flag='',
        verification=('Read directly from the official DGBAS workbook (%s, Table 2, downloaded from %s); identical '
                      'to Table 2 of the published yearbook (%s, PDF pages %d-%d, printed pages 4-5, downloaded from '
                      '%s). Retrieved %s.'
                      % (TW_XLSX, TW_XLSX_URL, TW_PDF, TW_PAGES[0], TW_PAGES[1], TW_PDF_URL, COLLECTED)),
        collected_on=COLLECTED))
    ext_rows.append(([u'%d (民國%d年)' % (y, y - 1911), '%.2f' % n[2], '%.2f' % n[7], '%.2f' % pub,
                      '%.2f' % n[9], '%.2f' % n[10], r.OBS_VALUE + ('' if r.OBS_STATUS == 'R' else ' (imputed)')],
                     True))
print('Taiwan: %d values read from the DGBAS workbook, identical to the yearbook'
      % sum(1 for r in rows if r['iso3'] == 'TWN'))

out = L.write_stage('unemployment_cells.csv', rows)
print('cells written: %d (%d countries x %d years)' % (len(out), len(ISO), len(YEARS)))
assert len(out) == len(ISO) * len(YEARS)

# the three sources side by side, cell by cell, for the record
cmp_rows = []
for iso in ISO:
    for y in YEARS:
        r, w = ilo.get((iso, y)), wb.get((iso, y))
        used = next(c for c in rows if c['iso3'] == iso and c['year'] == y)
        cmp_rows.append(dict(
            iso3=iso, year=y, ilo_value=r.OBS_VALUE, ilo_obs_status=r.OBS_STATUS or '(none: imputed)',
            ilo_lower_bound=r.LOWER_BOUND, ilo_upper_bound=r.UPPER_BOUND,
            world_bank_value='' if w is None else repr(w),
            world_bank_rounds_to_ilo='' if w is None else ('yes' if round(w, 3) == float(r.OBS_VALUE) else 'no'),
            dgbas_value=('%.2f' % used['value']) if iso == 'TWN' else '',
            used='DGBAS' if iso == 'TWN' else 'ILO', used_value=repr(float(used['value']))))
pd.DataFrame(cmp_rows).to_csv(os.path.join(D, 'unemployment_source_comparison.csv'), index=False,
                              encoding='utf-8-sig')

# ------------------------------------------------------------------ mirrors, extracts, publisher page
if FETCH:
    mirrors = []
    g_ilo = {iso: {y: ilo[(iso, y)].OBS_VALUE + ('' if ilo[(iso, y)].OBS_STATUS == 'R' else ' i') for y in YEARS}
             for iso in ISO}
    base = os.path.join(API, 'MIRROR__' + os.path.splitext(ILO_FILE)[0])
    made = L.grid_mirror('ILO modelled estimates (Nov. 2025) - unemployment rate, total, age 15+ (%)',
                         'Query: %s. Retrieved %s. %d panel countries, values as published (three decimals).'
                         % (ILO_URL, COLLECTED, len(g_ilo)), g_ilo, YEARS, base, fmt=lambda v: v,
                         note='"i" marks a value without the ILO observation status R (Real value): an ILO '
                              'imputation, not based on national data. Taiwan is shown for comparison; the panel uses '
                              'DGBAS for Taiwan. Rendered %s from the archived response beside this file.' % COLLECTED)
    mirrors.append(dict(scope='API', folder='evidence/api', source_file=ILO_FILE,
                        mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    g_wb = {iso: {y: wb[(iso, y)] for y in YEARS if (iso, y) in wb} for iso in ISO if iso != 'TWN'}
    base = os.path.join(API, 'MIRROR__' + os.path.splitext(WB_FILE)[0])
    made = L.grid_mirror('World Bank WDI SL.UEM.TOTL.ZS - unemployment, total (% of labor force) (modeled ILO estimate)',
                         'Query: api.worldbank.org country/all/indicator/SL.UEM.TOTL.ZS, 2001-2022. Retrieved %s. %d '
                         'panel countries; values shown to 3 decimals, the payload holds full precision.'
                         % (COLLECTED, len(g_wb)), g_wb, YEARS, base, fmt=lambda v: '%.3f' % v)
    mirrors.append(dict(scope='API', folder='evidence/api', source_file=WB_FILE,
                        mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    for rel, title, header, trows, src in (
            (WB_META, 'World Bank indicator metadata - SL.UEM.TOTL.ZS', ['Field', 'Value'],
             [(['id', meta['id']], False), (['name', meta['name']], True), (['sourceNote', meta['sourceNote']], True),
              (['sourceOrganization', meta['sourceOrganization']], True), (['source', meta['source']['value']], False)],
             WB_META_URL),
            (ILO_DF, 'ILO SDMX dataflow DF_UNE_2EAP_SEX_AGE_RT - name, last update and description', ['Field', 'Value'],
             [(['Name (en)', ilo_name], True), (['LAST_UPDATE annotation (day/month/year)',
                                                re.search(r'<common:AnnotationTitle>([^<]+)</common:AnnotationTitle>\s*'
                                                          r'<common:AnnotationType>LAST_UPDATE', dft).group(1)], True),
              (['Description (en)', ilo_desc], True)], ILO_DF_URL),
            (ILO_CL, 'ILO SDMX code list CL_OBS_STATUS - observation status', ['Code', 'Name (en)'],
             [([k, v], k == 'R') for k, v in codes.items()], ILO_CL_URL)):
        mb = os.path.join(API, 'MIRROR__' + os.path.splitext(rel)[0])
        L.table_extract(title, header, trows, mb, 'Archived response %s from %s, retrieved %s.' % (rel, src, COLLECTED),
                        note='The fields of the archived response, unchanged (markup removed from the description). '
                             'Rendered %s.' % COLLECTED)
        mirrors.append(dict(scope='API', folder='evidence/api', source_file=rel,
                            mirrors='%s;%s' % (os.path.basename(mb + '.pdf'), os.path.basename(mb + '.png')),
                            n_mirrors=2))
    # the yearbook: the two pages of Table 2 and the page of definitions
    made = []
    for pg in TW_PAGES + (TW_DEF_PAGE,):
        made += L.page_extract(os.path.join(TWN_DIR, TW_PDF), pg,
                               os.path.join(TWN_DIR, 'MIRROR__%s__p%d' % (os.path.splitext(TW_PDF)[0], pg)))
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_PDF,
                        mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    # the workbook: Table 2 as it stands, every row
    def fmt(v):
        return '%.2f' % v if isinstance(v, (int, float)) else str(v).strip()
    xrows = [([fmt(v) for v in vals[:12]],
              bool(YR.match(str(vals[0]))) and 2001 <= int(YR.match(str(vals[0])).group(2)) <= 2022)
             for vals in allrows]
    xb = os.path.join(TWN_DIR, 'MIRROR__' + os.path.splitext(TW_XLSX)[0])
    L.table_extract('DGBAS table2.xlsx - 表2 歷年人力資源調查重要指標 (all rows)',
                    ['年月別', '15+民間人口/總人口', '勞動力/總人口', '勞參率 總計', '勞參率 男', '勞參率 女',
                     '就業者/總人口', '就業者/15+民間人口', '就業者/勞動力', '失業率 總計', '失業率 男', '失業率 女'],
                    xrows, xb, 'Archived workbook %s, downloaded from %s on %s. Unit: %%.' % (TW_XLSX, TW_XLSX_URL, COLLECTED),
                    note='Every data row of the sheet as stored in the workbook; the highlighted rows (2001-2022) are '
                         'the ones the panel uses (column 失業率 總計). Rendered %s.' % COLLECTED)
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_XLSX,
                        mirrors='%s;%s' % (os.path.basename(xb + '.pdf'), os.path.basename(xb + '.png')), n_mirrors=2))
    # the page that links the files, rendered from a copy that names the page's own address as its base
    hb = os.path.join(TWN_DIR, 'MIRROR__' + os.path.splitext(TW_HTML)[0])
    raw = open(os.path.join(TWN_DIR, TW_HTML), encoding='utf-8', errors='replace').read()
    tmp = hb + '.tmp.html'
    open(tmp, 'w', encoding='utf-8').write(re.sub(r'(<head[^>]*>)', r'\1<base href="%s">' % TW_PAGE_URL,
                                                  raw, count=1, flags=re.I))
    made = L.chrome_render('file:///' + tmp.replace(chr(92), '/'), pdf=hb + '.pdf', png=hb + '.png')
    os.remove(tmp)
    mirrors.append(dict(scope='TWN', folder='evidence/countries/TWN', source_file=TW_HTML,
                        mirrors=';'.join(os.path.basename(m) for m in made), n_mirrors=len(made)))
    sm = os.path.join(D, 'source_mirrors.csv')
    old = pd.read_csv(sm)
    old = old[~(old.source_file.isin([m['source_file'] for m in mirrors]))]
    pd.concat([old, pd.DataFrame(mirrors)], ignore_index=True).to_csv(sm, index=False, encoding='utf-8-sig')

    # snapshots attached to the yearbook-page URL, which is the panel's source URL for Taiwan
    epdf, epng, made = L.table_extract(
        'Taiwan - unemployment rate, 2001-2022 (DGBAS Manpower Survey, Yearbook 2025 Table 2)',
        ['Year', 'Labour force participation rate', 'Employed / labour force', 'Unemployment rate (used)',
         'Unemployment rate, male', 'Unemployment rate, female', 'ILO modelled estimate (not used)'], ext_rows,
        os.path.join(TWN_DIR, TW_EXTRACT),
        'Source: %s. Workbook %s and yearbook %s (PDF pages %d-%d), both linked from %s. Unit: %%.'
        % (TW_NAME, TW_XLSX, TW_PDF, TW_PAGES[0], TW_PAGES[1], TW_PAGE_URL),
        note='Rows read from the archived workbook; the yearbook prints the same figures. The fourth column is the '
             'value used; the last column is the ILO\'s value for comparison (2021 and 2022 are ILO imputations). '
             'Extract prepared %s.' % COLLECTED)
    L.add_snapshot_row('TWN', VAR, TW_PAGE_URL, epdf, epng, TW_XLSX,
                       'Rows 2001-2022 of DGBAS Table 2 (the workbook; the yearbook pages are MIRROR__%s__p20/p21, the '
                       'definitions MIRROR__%s__p%d).' % (os.path.splitext(TW_PDF)[0], os.path.splitext(TW_PDF)[0],
                                                          TW_DEF_PAGE),
                       captured=COLLECTED)
    # the live site answers a headless browser with "Web Page Blocked!" (curl gets the page), so the snapshot is
    # rendered from the copy retrieved on the collection date, like the mirror above
    sp_pdf, sp_png = os.path.join(TWN_DIR, TW_PAGESHOT + '.pdf'), os.path.join(TWN_DIR, TW_PAGESHOT + '.png')
    for f in (sp_pdf, sp_png):
        if os.path.exists(f):
            os.remove(f)
    L.add_snapshot_row('TWN', VAR, TW_PAGE_URL, hb + '.pdf', hb + '.png', TW_HTML,
                       'The DGBAS statistics page 114年人力資源調查統計, which links the yearbook (PDF) and its tables '
                       '(表2 歷年人力資源調查重要指標, EXCEL), rendered from the copy retrieved %s: the live site blocks '
                       'headless browsers.' % COLLECTED, status='rendered_from_archived_html', captured=COLLECTED)
    ws_ = pd.read_csv(os.path.join(D, 'web_snapshots.csv'))
    ws_[~ws_.png_screenshot.astype(str).eq(TW_PAGESHOT + '.png')].to_csv(
        os.path.join(D, 'web_snapshots.csv'), index=False, encoding='utf-8-sig')

    pubrows = []
    for p in PUBLISHER_PAGES:
        b = os.path.join(PUB, p['key'])
        made = L.chrome_render(p['page_url'], pdf=b + '.pdf', png=b + '.png', wait_ms=20000)
        pubrows.append(dict(p, status='captured' if len(made) == 2 else 'partial',
                            pdf='evidence/api/publisher_pages/%s.pdf' % p['key'] if b + '.pdf' in made else '',
                            png='evidence/api/publisher_pages/%s.png' % p['key'] if b + '.png' in made else '',
                            pdf_bytes=os.path.getsize(b + '.pdf') if b + '.pdf' in made else 0,
                            png_bytes=os.path.getsize(b + '.png') if b + '.png' in made else 0,
                            captured=COLLECTED, note=''))
        print('publisher page %-30s %s' % (p['key'], [os.path.basename(m) for m in made]))
    upsert(os.path.join(D, 'api_publisher_snapshots.csv'), pubrows, 'key')
